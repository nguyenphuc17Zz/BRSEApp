import os
import docx
from docx.shared import Pt, Inches, RGBColor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import pptx
from pptx.util import Inches as PptxInches, Pt as PptxPt
import fitz

OUTPUT_DIR = r"E:\AutomationTranslate\sample_documents"
os.makedirs(OUTPUT_DIR, exist_ok=True)

def create_sample_docx():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading("システム認証仕様書 (System Authentication Spec)", level=0)
    
    # Section 1
    doc.add_heading("1. 概要 (Overview)", level=1)
    p1 = doc.add_paragraph("本仕様書は、次期バンキングシステムにおけるOAuth2 Authorization Code Flowを用いたユーザー認証機能の基本仕様を定義するものです。")
    
    # Section 2
    doc.add_heading("2. セキュリティ要件 (Security Requirements)", level=1)
    p2 = doc.add_paragraph()
    p2.add_run("・セッション有効期限は無操作後30分とします。\n")
    p2.add_run("・パスワード誤入力が連続5回発生した場合、アカウントを一時ロックします。\n")
    p2.add_run("・トークン更新時はリフレッシュトークンによるPKCE検証を必須とします。")
    
    # Section 3 Table
    doc.add_heading("3. API一覧 (API List)", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "エンドポイント (Endpoint)"
    hdr_cells[1].text = "メソッド (Method)"
    hdr_cells[2].text = "概要 (Summary)"
    
    apis = [
        ("/api/v1/auth/login", "POST", "ユーザーログイン及びトークン発行処理"),
        ("/api/v1/auth/refresh", "POST", "アクセストークンの再発行処理"),
        ("/api/v1/auth/logout", "POST", "セッション破棄及びログアウト処理")
    ]
    for ep, method, desc in apis:
        row_cells = table.add_row().cells
        row_cells[0].text = ep
        row_cells[1].text = method
        row_cells[2].text = desc

    path = os.path.join(OUTPUT_DIR, "01_Specification_Auth.docx")
    doc.save(path)
    print(f"Created {path}")

def create_sample_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工数見積 (Estimate)"
    
    # Headers
    headers = ["機能名 (Feature)", "担当 (Role)", "工数[人日] (Effort)", "単価[万円] (Rate)", "金額[万円] (Subtotal)"]
    ws.append(headers)
    
    # Data rows
    ws.append(["認証API実装", "Backend Lead", 10, 50, "=C2*D2"])
    ws.append(["ログイン画面UI開発", "Frontend Dev", 8, 45, "=C3*D3"])
    ws.append(["単体テストおよび結合テスト", "QA Engineer", 12, 40, "=C4*D4"])
    ws.append(["セキュリティ監査対応", "BrSE", 5, 55, "=C5*D5"])
    
    # Total row with formulas
    ws.append(["合計 (Total)", "", "=SUM(C2:C5)", "", "=SUM(E2:E5)"])
    
    # Style header
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        
    path = os.path.join(OUTPUT_DIR, "02_Project_Estimates.xlsx")
    wb.save(path)
    print(f"Created {path}")

def create_sample_pptx():
    prs = pptx.Presentation()
    
    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    slide1.shapes.title.text = "クラウド移行アーキテクチャ計画"
    slide1.placeholders[1].text = "2026年度 次期コアシステム刷新プロジェクト\n作成者: BrSEチーム"
    
    # Slide 2: Bullet points + Notes
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    slide2.shapes.title.text = "システム移行方針と技術選定"
    tf = slide2.shapes.placeholders[1].text_frame
    tf.text = "マイクロサービスアーキテクチャへの段階的移行"
    p = tf.add_paragraph()
    p.text = "データベースのマルチAZ冗長化による高可用性の確保"
    p2 = tf.add_paragraph()
    p2.text = "API Gateway経由でのOAuth2トークン集中検証"
    
    # Speaker notes
    notes_slide = slide2.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "【発表用メモ】クライアント経営陣へ説明する際、ダウンタイムゼロで安全に切り替える手順を強調して説明してください。"
    
    path = os.path.join(OUTPUT_DIR, "03_Architecture_Review.pptx")
    prs.save(path)
    print(f"Created {path}")

def create_sample_pdf():
    doc = fitz.open()
    page = doc.new_page(width=595, height=842) # A4
    
    rect_title = fitz.Rect(50, 50, 545, 90)
    page.insert_textbox(rect_title, "Information Security Policy (情報セキュリティ方針)", fontsize=18, fontname="helv")
    
    rect_body = fitz.Rect(50, 110, 545, 400)
    content = """1. Password & Access Control:
- Minimum length: 12 characters with mixed case and symbols.
- Multi-factor authentication (MFA) is strictly required for production admin access.

2. Incident Reporting Procedure:
- Any security anomaly or HTTP 500 spike must be escalated to the BrSE within 15 minutes.
- Root cause analysis (RCA) must be finalized within 24 hours of mitigation.

3. Data Classification:
- Customer PII must be encrypted at rest with AES-256 and masked in logs.
"""
    page.insert_textbox(rect_body, content, fontsize=11, fontname="helv")
    
    path = os.path.join(OUTPUT_DIR, "04_Security_Policy.pdf")
    doc.save(path)
    print(f"Created {path}")

if __name__ == "__main__":
    create_sample_docx()
    create_sample_xlsx()
    create_sample_pptx()
    create_sample_pdf()
