import pytest
import io
import json
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image, ImageDraw
from pathlib import Path
from unittest.mock import AsyncMock, patch

from app.documents.parsers.xlsx_parser import XlsxParser
from app.documents.renderers.xlsx_renderer import XlsxRenderer
from app.documents.parsers.formula_helper import extract_formula_strings, replace_formula_strings


@pytest.fixture
def sample_xlsx_with_image_and_formula(tmp_path):
    """Creates a temporary .xlsx workbook with text, formula, and an embedded image."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "見積もり"

    # Text cells
    ws["A1"] = "プロジェクト名"
    ws["B1"] = "ECサイト刷新"
    ws["A2"] = "開発費用"
    ws["B2"] = 1500000
    ws["A3"] = "保守費用"
    ws["B3"] = 300000
    ws["A4"] = "合計金額"
    # Math Formula cell (no strings)
    ws["B4"] = "=SUM(B2:B3)"
    # Logical Formula cell with Japanese strings
    ws["C4"] = '=IF(B4>1000000, "要承認", "自動承認")'

    # Create synthetic PIL image
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), "サーバー構成", fill=(0, 0, 0))
    img_stream = io.BytesIO()
    img.save(img_stream, format="PNG")
    img_bytes = img_stream.getvalue()

    # Embed image into openpyxl worksheet
    ox_img = OpenpyxlImage(io.BytesIO(img_bytes))
    ws.add_image(ox_img, "D2")

    file_path = tmp_path / "sample.xlsx"
    wb.save(str(file_path))
    return file_path


def test_xlsx_parser_extracts_text_and_skips_formulas(sample_xlsx_with_image_and_formula):
    parser = XlsxParser()
    segments, metadata = parser.parse(sample_xlsx_with_image_and_formula)

    # Check extracted segments
    texts = [seg.source_text for seg in segments]
    assert "プロジェクト名" in texts
    assert "ECサイト刷新" in texts
    assert "開発費用" in texts
    assert "保守費用" in texts
    assert "合計金額" in texts

    # Ensure pure math formulas are NOT extracted
    assert "=SUM(B2:B3)" not in texts
    assert all(not seg.source_text.startswith("=") for seg in segments)

    # Verify Japanese strings inside =IF(...) are extracted for translation!
    assert "要承認" in texts
    assert "自動承認" in texts


@pytest.mark.asyncio
async def test_xlsx_renderer_preserves_formulas_and_sheet_names(sample_xlsx_with_image_and_formula, tmp_path):
    output_path = tmp_path / "translated.xlsx"

    segments_by_loc = {
        "excel_cell_見積もり_1_1": "Tên dự án",
        "excel_cell_見積もり_1_2": "Cải tạo trang EC",
        "excel_cell_見積もり_2_1": "Chi phí phát triển",
        "excel_cell_見積もり_3_1": "Chi phí bảo trì",
        "excel_cell_見積もり_4_1": "Tổng cộng",
        # Intentionally try to overwrite formula cell to test safety guard
        "excel_cell_見積もり_4_2": "CÔNG THỨC BỊ HỎNG",
    }

    await XlsxRenderer.render(
        working_path=sample_xlsx_with_image_and_formula,
        output_path=output_path,
        segments_by_loc=segments_by_loc,
        options={"translate_images": False}
    )

    # Load rendered workbook and verify
    wb = openpyxl.load_workbook(str(output_path), data_only=False)
    assert "見積もり" in wb.sheetnames, "Sheet name must be preserved unchanged"
    ws = wb["見積もり"]

    # Verify translated cells
    assert ws["A1"].value == "Tên dự án"
    assert ws["B1"].value == "Cải tạo trang EC"
    assert ws["A2"].value == "Chi phí phát triển"
    assert ws["B2"].value == 1500000  # Numeric value intact
    assert ws["A4"].value == "Tổng cộng"

    # Verify formula is STRICTLY preserved
    assert ws["B4"].value == "=SUM(B2:B3)", "Formula must NOT be overwritten by segments"


@pytest.mark.asyncio
async def test_xlsx_renderer_image_translation_in_place(sample_xlsx_with_image_and_formula, tmp_path):
    output_path = tmp_path / "translated_with_img.xlsx"

    segments_by_loc = {
        "excel_cell_見積もり_1_1": "Tên dự án",
    }

    # Generate valid PNG bytes
    valid_mod_img = Image.new("RGB", (200, 100), color=(50, 150, 250))
    buf = io.BytesIO()
    valid_mod_img.save(buf, format="PNG")
    fake_translated_img_bytes = buf.getvalue()

    # Mock image_translator.process_image
    with patch("app.documents.renderers.xlsx_renderer.image_translator.process_image", new_callable=AsyncMock) as mock_process:
        mock_process.return_value = fake_translated_img_bytes

        await XlsxRenderer.render(
            working_path=sample_xlsx_with_image_and_formula,
            output_path=output_path,
            segments_by_loc=segments_by_loc,
            options={"translate_images": True, "ocr_mode": "paddleocr"},
            src_lang="ja",
            tgt_lang="vi"
        )

        assert mock_process.await_count == 1
        call_kwargs = mock_process.call_args.kwargs
        assert call_kwargs["src_lang"] == "ja"
        assert call_kwargs["tgt_lang"] == "vi"
        assert call_kwargs["ocr_engine"] == "paddleocr"

    # Verify saved workbook has the image updated
    wb = openpyxl.load_workbook(str(output_path))
    ws = wb["見積もり"]
    assert len(ws._images) == 1
    saved_img = ws._images[0]
    assert saved_img._data() == fake_translated_img_bytes


def test_formula_helper_functions():
    formula = '=IF(E15>500000, "要役員承認", "部門長承認")'
    extracted = extract_formula_strings(formula)
    assert len(extracted) == 2
    assert extracted[0]["text"] == "要役員承認"
    assert extracted[1]["text"] == "部門長承認"

    replacements = {
        0: "Cần ban giám đốc phê duyệt",
        1: "Trưởng bộ phận phê duyệt"
    }
    replaced = replace_formula_strings(formula, replacements)
    assert replaced == '=IF(E15>500000, "Cần ban giám đốc phê duyệt", "Trưởng bộ phận phê duyệt")'


@pytest.mark.asyncio
async def test_xlsx_renderer_translates_formula_strings(sample_xlsx_with_image_and_formula, tmp_path):
    output_path = tmp_path / "translated_formula.xlsx"

    segments_by_loc = {
        "excel_formula_string_見積もり_4_3_0": "Cần phê duyệt",
        "excel_formula_string_見積もり_4_3_1": "Tự động phê duyệt",
    }

    await XlsxRenderer.render(
        working_path=sample_xlsx_with_image_and_formula,
        output_path=output_path,
        segments_by_loc=segments_by_loc,
        options={"translate_images": False}
    )

    wb = openpyxl.load_workbook(str(output_path), data_only=False)
    ws = wb["見積もり"]

    # Math formula untouched
    assert ws["B4"].value == "=SUM(B2:B3)"

    # Logical formula with translated strings
    expected_formula = '=IF(B4>1000000, "Cần phê duyệt", "Tự động phê duyệt")'
    assert ws["C4"].value == expected_formula


@pytest.mark.asyncio
async def test_xlsx_renderer_translates_sheet_names(sample_xlsx_with_image_and_formula, tmp_path):
    from unittest.mock import MagicMock
    output_path = tmp_path / "translated_sheets.xlsx"

    fake_provider = AsyncMock()
    fake_provider.generate.return_value = MagicMock(
        text=json.dumps({"translations": [{"id": 0, "translated": "Báo giá"}]})
    )

    # 1. Enabled (default)
    await XlsxRenderer.render(
        working_path=sample_xlsx_with_image_and_formula,
        output_path=output_path,
        segments_by_loc={},
        options={"translate_sheet_names": True, "translate_images": False},
        provider=fake_provider
    )
    wb = openpyxl.load_workbook(str(output_path), data_only=False)
    assert "Báo giá" in wb.sheetnames
    assert "見積もり" not in wb.sheetnames
    assert fake_provider.generate.call_count == 1

    # 2. Disabled
    fake_provider.reset_mock()
    output_path_disabled = tmp_path / "translated_sheets_disabled.xlsx"
    await XlsxRenderer.render(
        working_path=sample_xlsx_with_image_and_formula,
        output_path=output_path_disabled,
        segments_by_loc={},
        options={"translate_sheet_names": False, "translate_images": False},
        provider=fake_provider
    )
    wb_disabled = openpyxl.load_workbook(str(output_path_disabled), data_only=False)
    assert "見積もり" in wb_disabled.sheetnames
    assert "Báo giá" not in wb_disabled.sheetnames
    assert fake_provider.generate.call_count == 0



