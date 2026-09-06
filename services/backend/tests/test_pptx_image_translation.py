import pytest
import io
from pathlib import Path
from unittest.mock import AsyncMock, patch
import pptx
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw

from app.documents.parsers.pptx_parser import PptxParser
from app.documents.renderers.pptx_renderer import PptxRenderer


@pytest.fixture
def sample_pptx_with_image(tmp_path):
    """Creates a sample presentation with title, table, notes, and an embedded image."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # 1. Textbox
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(1))
    p = tx_box.text_frame.paragraphs[0]
    p.text = "クラウドアーキテクチャ提案書"

    # 2. Table
    table_shape = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1.8), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "サービス"
    table.cell(0, 1).text = "用途"
    table.cell(1, 0).text = "ECS Fargate"
    table.cell(1, 1).text = "コンテナ実行基盤"

    # 3. Speaker notes
    slide.notes_slide.notes_text_frame.text = "本スライドでは全体構成を説明します。"

    # 4. Embedded image
    img = Image.new("RGB", (250, 120), color=(240, 248, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, 100, 50], fill=(41, 128, 185))
    draw.text((20, 25), "AWS WAF", fill=(255, 255, 255))
    draw.rectangle([120, 10, 230, 50], fill=(39, 174, 96))
    draw.text((130, 25), "ALB", fill=(255, 255, 255))

    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_bytes = img_buf.getvalue()

    slide.shapes.add_picture(io.BytesIO(img_bytes), Inches(5.2), Inches(1.8), Inches(3.8), Inches(2))

    pptx_path = tmp_path / "sample_arch.pptx"
    prs.save(str(pptx_path))
    return pptx_path


def test_pptx_parser_extracts_shapes_and_notes(sample_pptx_with_image):
    parser = PptxParser()
    segments, metadata = parser.parse(sample_pptx_with_image)

    texts = [seg.source_text for seg in segments]
    assert any("クラウドアーキテクチャ提案書" in t for t in texts)
    assert any("ECS Fargate" in t for t in texts)
    assert any("コンテナ実行基盤" in t for t in texts)
    assert any("本スライドでは全体構成" in t for t in texts)


@pytest.mark.asyncio
async def test_pptx_renderer_text_and_image_translation(sample_pptx_with_image, tmp_path):
    parser = PptxParser()
    segments, _ = parser.parse(sample_pptx_with_image)

    segments_by_loc = {}
    for s in segments:
        loc = s.location
        loc_type = loc.get("type")
        if loc_type == "pptx_shape_text":
            key = f"pptx_shape_text_{loc['slide_index']}_{loc['shape_index']}_{loc['paragraph_index']}"
        elif loc_type == "pptx_table_cell":
            key = f"pptx_table_cell_{loc['slide_index']}_{loc['shape_index']}_{loc['row_index']}_{loc['col_index']}"
        elif loc_type == "pptx_speaker_notes":
            key = f"pptx_speaker_notes_{loc['slide_index']}_{loc['paragraph_index']}"
        segments_by_loc[key] = f"Dịch: {s.source_text}"

    # Mock translated image bytes
    mock_translated_img = Image.new("RGB", (250, 120), color=(200, 255, 200))
    m_buf = io.BytesIO()
    mock_translated_img.save(m_buf, format="PNG")
    mock_bytes = m_buf.getvalue()

    output_pptx = tmp_path / "output_translated.pptx"

    with patch("app.documents.ocr.image_translator.image_translator.process_image", new=AsyncMock(return_value=mock_bytes)):
        await PptxRenderer.render(
            working_path=sample_pptx_with_image,
            output_path=output_pptx,
            segments_by_loc=segments_by_loc,
            options={"translate_images": True, "ocr_mode": "auto"},
            src_lang="ja",
            tgt_lang="vi"
        )

    assert output_pptx.exists()

    # Reopen to verify
    out_prs = pptx.Presentation(str(output_pptx))
    slide = out_prs.slides[0]

    # Check text
    all_texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_texts.extend(p.text for p in shape.text_frame.paragraphs)
        elif shape.has_table:
            for row in shape.table.rows:
                all_texts.extend(cell.text for cell in row.cells)
    if slide.has_notes_slide:
        all_texts.extend(p.text for p in slide.notes_slide.notes_text_frame.paragraphs)

    assert any("Dịch: クラウドアーキテクチャ提案書" in t for t in all_texts)
    assert any("Dịch: ECS Fargate" in t for t in all_texts)
    assert any("Dịch: 本スライドでは全体構成" in t for t in all_texts)

    # Check image was replaced in-place
    picture_shape = [s for s in slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE][0]
    rId = picture_shape._element.blipFill.blip.rEmbed
    img_part = slide.part.related_part(rId)
    assert img_part.blob == mock_bytes
