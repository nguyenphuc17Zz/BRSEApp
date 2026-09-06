import asyncio
import json
from pathlib import Path
from typing import Dict, List, Optional, Any
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import async_session_maker
from app.core.logging import logger
from app.documents.models import DocumentJob, DocumentSegment, DocumentFile, DocumentIssue
from app.documents.storage import document_storage
from app.documents.parsers.docx_parser import DocxParser
from app.documents.parsers.xlsx_parser import XlsxParser
from app.documents.parsers.pptx_parser import PptxParser
from app.documents.parsers.pdf_parser import PdfParser
from app.documents.renderers.docx_renderer import DocxRenderer
from app.documents.renderers.xlsx_renderer import XlsxRenderer
from app.documents.renderers.pptx_renderer import PptxRenderer
from app.documents.renderers.pdf_renderer import PdfRenderer
from app.documents.qa import DocumentQAChecker
from app.documents.segmenter import TokenProtector
from app.documents.ocr.image_translator import image_translator
from app.engine.context import context_engine
from app.engine.prompt import PromptBuilder
from app.engine.pipeline import clean_json_response
from app.providers.registry import provider_registry

class JobManager:
    """Manages asynchronous document translation job lifecycles and background workers."""

    def __init__(self):
        self._active_tasks: Dict[str, asyncio.Task] = {}
        self._cancel_flags: Dict[str, bool] = {}
        self._pause_flags: Dict[str, bool] = {}

    def start_job(self, job_id: str):
        """Starts document translation in a background task."""
        if job_id in self._active_tasks and not self._active_tasks[job_id].done():
            logger.warning(f"Job {job_id} is already running.")
            return

        self._cancel_flags[job_id] = False
        self._pause_flags[job_id] = False
        task = asyncio.create_task(self._run_job_pipeline(job_id))
        self._active_tasks[job_id] = task

    def pause_job(self, job_id: str):
        self._pause_flags[job_id] = True

    def resume_job(self, job_id: str):
        self._pause_flags[job_id] = False
        self.start_job(job_id)

    def cancel_job(self, job_id: str):
        self._cancel_flags[job_id] = True
        if job_id in self._active_tasks:
            self._active_tasks[job_id].cancel()

    async def _run_job_pipeline(self, job_id: str):
        """Executes full document translation stages: analyzing -> segmenting -> translating -> QA -> rendering."""
        try:
            async with async_session_maker() as db:
                job = (await db.execute(select(DocumentJob).where(DocumentJob.id == job_id))).scalar_one_or_none()
                if not job:
                    logger.error(f"Job {job_id} not found.")
                    return

                doc_file = (await db.execute(select(DocumentFile).where(DocumentFile.id == job.document_id))).scalar_one_or_none()
                if not doc_file:
                    job.status = "failed"
                    job.error_message = "Document file not found."
                    await db.commit()
                    return

                options = json.loads(job.options_json or "{}")

                # Stage 1: Analyzing & Segmenting (if not already segmented)
                seg_chk = (await db.execute(select(DocumentSegment).where(DocumentSegment.job_id == job_id))).scalars().first()
                if not seg_chk:
                    job.status = "segmenting"
                    job.current_stage = "Extracting document structure and text segments..."
                    await db.commit()

                    file_path = Path(doc_file.original_path)
                    parser = self._get_parser(doc_file.file_type)
                    parsed_segments, meta = parser.parse(file_path, options)

                    # Persist segments in SQLite
                    for ps in parsed_segments:
                        seg = DocumentSegment(
                            job_id=job.id,
                            segment_index=ps.segment_index,
                            location_json=json.dumps(ps.location, ensure_ascii=False),
                            source_text=ps.source_text,
                            protected_tokens_json=json.dumps(ps.protected_tokens, ensure_ascii=False),
                            context_hint=ps.context_hint,
                            status="pending"
                        )
                        db.add(seg)

                    job.total_segments = len(parsed_segments)
                    await db.commit()

                # Stage 2: Translating in Batches (with deduplication cache)
                job.status = "translating"
                await db.commit()

                pending_segs_res = await db.execute(
                    select(DocumentSegment)
                    .where(DocumentSegment.job_id == job_id, DocumentSegment.status == "pending")
                    .order_by(DocumentSegment.segment_index.asc())
                )
                pending_segments = pending_segs_res.scalars().all()

                provider = provider_registry.get_provider(job.provider) or provider_registry.get_provider("gemini")
                translation_cache: Dict[str, str] = {}
                max_batch_chars = 1200
                max_batch_items = 8 if (job.provider or "").lower() == "groq" else 12

                batches: List[List[DocumentSegment]] = []
                current_batch: List[DocumentSegment] = []
                current_chars = 0

                for seg in pending_segments:
                    seg_len = len(seg.source_text or "")
                    if current_batch and (current_chars + seg_len > max_batch_chars or len(current_batch) >= max_batch_items):
                        batches.append(current_batch)
                        current_batch = [seg]
                        current_chars = seg_len
                    else:
                        current_batch.append(seg)
                        current_chars += seg_len

                if current_batch:
                    batches.append(current_batch)

                processed_so_far = 0
                for b_idx, batch in enumerate(batches):
                    if self._cancel_flags.get(job_id):
                        job.status = "cancelled"
                        await db.commit()
                        return

                    if self._pause_flags.get(job_id):
                        job.status = "paused"
                        await db.commit()
                        return

                    processed_so_far += len(batch)
                    job.current_stage = f"Translating batch {b_idx + 1}/{len(batches)} ({processed_so_far}/{job.total_segments} items)..."
                    await db.commit()

                    await self._translate_batch(
                        db=db,
                        batch=batch,
                        job=job,
                        provider=provider,
                        translation_cache=translation_cache
                    )

                    # Update progress
                    completed_cnt = (await db.execute(
                        select(DocumentSegment).where(DocumentSegment.job_id == job_id, DocumentSegment.status.in_(["translated", "user_edited"]))
                    )).scalars().all()
                    failed_cnt = (await db.execute(
                        select(DocumentSegment).where(DocumentSegment.job_id == job_id, DocumentSegment.status == "failed")
                    )).scalars().all()
                    job.completed_segments = len(completed_cnt)
                    job.failed_segments = len(failed_cnt)
                    processed_total = job.completed_segments + job.failed_segments
                    job.progress_percent = round((processed_total / max(1, job.total_segments)) * 100, 1)
                    await db.commit()

                    # Adaptive pacing between batches to respect OTPM limits
                    if b_idx < len(batches) - 1:
                        pace_delay = 1.2 if (job.provider or "").lower() == "groq" else 0.5
                        await asyncio.sleep(pace_delay)

                # Check if all segments failed (prevent silent success)
                if job.completed_segments == 0 and job.failed_segments > 0:
                    job.status = "failed"
                    job.progress_percent = 100.0
                    clean_err = job.error_message or "Không thể kết nối hoặc mô hình AI quá tải/hết quota."
                    job.error_message = clean_err
                    job.current_stage = f"Thất bại: {clean_err}"
                    await db.commit()
                    logger.error(f"Document translation job {job_id} failed: 0 of {job.total_segments} segments translated.")
                    return

                # Stage 3: Document QA
                job.status = "qa"
                job.current_stage = "Running document QA verification..."
                await db.commit()

                all_segs_res = await db.execute(
                    select(DocumentSegment).where(DocumentSegment.job_id == job_id)
                )
                all_segs = all_segs_res.scalars().all()

                for seg in all_segs:
                    token_map = json.loads(seg.protected_tokens_json or "{}")
                    issues = DocumentQAChecker.check_segment(
                        job_id=job.id,
                        segment_id=seg.id,
                        source_text=seg.source_text,
                        translated_text=seg.translated_text or "",
                        location_text=seg.context_hint or f"Segment #{seg.segment_index + 1}",
                        protected_tokens_map=token_map
                    )
                    for iss in issues:
                        db.add(iss)

                await db.commit()

                # Stage 4: Rendering Output Document
                job.status = "rendering"
                job.current_stage = "Rendering translated document..."
                await db.commit()

                job_options = json.loads(job.options_json or "{}")
                custom_output_dir = job_options.get("custom_output_dir")
                custom_filename = job_options.get("target_filename") or job_options.get("custom_output_filename")

                working_copy = document_storage.create_working_copy(Path(doc_file.original_path), job_id)
                output_path = document_storage.get_output_path(
                    doc_file.filename,
                    job.target_language,
                    custom_dir=custom_output_dir,
                    custom_filename=custom_filename
                )

                await self._render_output_document(
                    file_type=doc_file.file_type,
                    working_copy=working_copy,
                    output_path=output_path,
                    segments=all_segs,
                    options=job_options,
                    src_lang=job.source_language,
                    tgt_lang=job.target_language,
                    provider=provider
                )

                document_storage.cleanup_working(working_copy)

                # Finalize Job
                job.progress_percent = 100.0
                job.output_filename = output_path.name
                job.output_path = str(output_path)

                if job.failed_segments > 0:
                    job.status = "partially_completed"
                    job.current_stage = f"Hoàn tất với cảnh báo ({job.completed_segments}/{job.total_segments} đoạn dịch thành công, {job.failed_segments} đoạn bị lỗi)."
                    logger.warning(f"Document translation job {job_id} finished with warnings: {job.failed_segments} segments failed.")
                else:
                    job.status = "completed"
                    job.current_stage = "Completed successfully. Output document ready for download."
                    logger.info(f"Document translation job {job_id} completed successfully!")
                await db.commit()

        except asyncio.CancelledError:
            logger.info(f"Job {job_id} cancelled.")
            async with async_session_maker() as db:
                job = (await db.execute(select(DocumentJob).where(DocumentJob.id == job_id))).scalar_one_or_none()
                if job and job.status not in ("completed", "partially_completed"):
                    job.status = "cancelled"
                    job.current_stage = "Đã hủy bởi người dùng"
                    await db.commit()
        except Exception as e:
            logger.error(f"Job {job_id} failed: {e}", exc_info=True)
            async with async_session_maker() as db:
                job = (await db.execute(select(DocumentJob).where(DocumentJob.id == job_id))).scalar_one_or_none()
                if job:
                    job.status = "failed"
                    job.error_message = str(e)
                    job.current_stage = f"Failed: {str(e)}"
                    await db.commit()

    async def _translate_batch(
        self,
        db: AsyncSession,
        batch: List[DocumentSegment],
        job: DocumentJob,
        provider: Any,
        translation_cache: Dict[str, str]
    ):
        """Translates a batch of segments, reusing cached translations for identical segments."""
        to_translate = []
        for seg in batch:
            if seg.source_text in translation_cache:
                seg.translated_text = translation_cache[seg.source_text]
                seg.status = "translated"
            else:
                to_translate.append(seg)

        if not to_translate:
            await db.commit()
            return

        # Build combined prompt for batch
        lang_names = {
            "ja": "Japanese",
            "vi": "Vietnamese",
            "en": "English"
        }
        src_lang_name = lang_names.get(job.source_language.lower(), job.source_language.upper())
        tgt_lang_name = lang_names.get(job.target_language.lower(), job.target_language.upper())

        example_translated = {
            "ja": "翻訳されたテキスト",
            "vi": "Văn bản đã được dịch",
            "en": "Translated text"
        }.get(job.target_language.lower(), f"Translated text in {tgt_lang_name}")

        items_payload = [{"id": s.segment_index, "text": s.source_text, "context": s.context_hint} for s in to_translate]
        prompt = f"""Translate the following document segments from {src_lang_name} to {tgt_lang_name} with tone/style: {job.style}.

CRITICAL REQUIREMENTS:
1. Every segment MUST be translated into {tgt_lang_name}. Do NOT output in {src_lang_name}.
2. Do NOT alter, remove, or translate any protected placeholders formatted as __PROTECTED_...__.
3. Preserve formatting, numbers, bullet styles, and tags identically.

Segments to translate (JSON):
{json.dumps(items_payload, ensure_ascii=False)}

Return JSON matching this exact structure:
{{
  "translations": [
    {{"id": 0, "translated": "{example_translated}"}}
  ]
}}"""
        system_instruction = (
            f"You are an expert technical and document translator specializing in {src_lang_name} to {tgt_lang_name}. "
            f"All translations must be strictly produced in {tgt_lang_name}. Preserve placeholders and structure."
        )

        max_batch_retries = 3
        last_batch_err = None

        for attempt_idx in range(max_batch_retries):
            try:
                resp_data = await provider.generate(
                    prompt=prompt,
                    system_instruction=system_instruction,
                    model=job.model,
                    temperature=0.2,
                    json_mode=True
                )
                parsed = clean_json_response(resp_data.text)
                trans_list = parsed.get("translations", [])
                if not trans_list:
                    raise RuntimeError(f"AI returned empty translation list for batch. Output snippet: {resp_data.text[:120]}")

                trans_map = {}
                for item in trans_list:
                    if "id" in item:
                        raw_id = item["id"]
                        val = item.get("translated", "")
                        trans_map[raw_id] = val
                        try:
                            trans_map[int(raw_id)] = val
                            trans_map[str(raw_id)] = val
                        except (ValueError, TypeError):
                            pass

                # Verify that at least one segment ID from to_translate was actually translated
                has_any_match = any(
                    s.segment_index in trans_map or str(s.segment_index) in trans_map
                    for s in to_translate
                )
                if not has_any_match:
                    raise RuntimeError("AI translation response did not match any segment IDs in current batch.")

                for seg in to_translate:
                    translated_raw = trans_map.get(seg.segment_index)
                    if translated_raw is None:
                        translated_raw = trans_map.get(str(seg.segment_index))

                    token_map = json.loads(seg.protected_tokens_json or "{}")

                    if translated_raw is not None and str(translated_raw).strip():
                        # Restore protected tokens
                        restored_text, missing = TokenProtector.restore_tokens(translated_raw, token_map)
                        seg.translated_text = restored_text
                        seg.status = "translated"
                        # Only cache if actually translated (prevent caching untranslated source text)
                        if restored_text.strip() != seg.source_text.strip():
                            translation_cache[seg.source_text] = restored_text
                    else:
                        # Mark this omitted segment as failed, keeping source text without polluting cache
                        restored_text, _ = TokenProtector.restore_tokens(seg.source_text, token_map)
                        seg.translated_text = restored_text
                        seg.status = "failed"
                        seg.error_message = "Segment omitted by AI translation model."

                last_batch_err = None
                break
            except Exception as e:
                last_batch_err = e
                if attempt_idx < max_batch_retries - 1:
                    wait_sec = 2.0 * (attempt_idx + 1)
                    logger.warning(f"Batch AI translation error on attempt {attempt_idx + 1}/{max_batch_retries}: {e}. Retrying in {wait_sec}s...")
                    await asyncio.sleep(wait_sec)
                else:
                    logger.warning(f"Batch AI translation failed after {max_batch_retries} attempts: {e}. Marking segments as failed.")

        if last_batch_err is not None:
            err_msg = str(last_batch_err)
            job.error_message = err_msg
            for seg in to_translate:
                token_map = json.loads(seg.protected_tokens_json or "{}")
                restored_text, _ = TokenProtector.restore_tokens(seg.source_text, token_map)
                seg.translated_text = restored_text
                seg.status = "failed"
                seg.error_message = err_msg

        await db.commit()

    def _get_parser(self, file_type: str):
        t = file_type.lower().replace(".", "")
        if t == "docx":
            return DocxParser()
        elif t == "xlsx":
            return XlsxParser()
        elif t == "pptx":
            return PptxParser()
        elif t == "pdf":
            return PdfParser()
        raise ValueError(f"Unsupported file type: {file_type}")

    async def _render_output_document(
        self,
        file_type: str,
        working_copy: Path,
        output_path: Path,
        segments: List[DocumentSegment],
        options: Optional[Dict[str, Any]] = None,
        src_lang: str = "ja",
        tgt_lang: str = "vi",
        provider: Any = None
    ):
        t = file_type.lower().replace(".", "")
        # Build lookup by location key
        segments_by_loc: Dict[str, str] = {}
        raw_segs = []

        for seg in segments:
            loc = json.loads(seg.location_json or "{}")
            text = seg.translated_text or seg.source_text
            raw_segs.append({"location": loc, "text": text})

            loc_type = loc.get("type")
            if loc_type == "paragraph":
                key = f"paragraph_{loc.get('p_index')}"
            elif loc_type == "table_cell":
                key = f"table_cell_{loc.get('t_index')}_{loc.get('r_index')}_{loc.get('c_index')}"
            elif loc_type == "excel_cell":
                key = f"excel_cell_{loc.get('sheet')}_{loc.get('row')}_{loc.get('column')}"
            elif loc_type == "excel_formula_string":
                key = f"excel_formula_string_{loc.get('sheet')}_{loc.get('row')}_{loc.get('column')}_{loc.get('str_index')}"
            elif loc_type == "pptx_shape_text":
                key = f"pptx_shape_text_{loc.get('slide_index')}_{loc.get('shape_index')}_{loc.get('paragraph_index')}"
            elif loc_type == "pptx_table_cell":
                key = f"pptx_table_cell_{loc.get('slide_index')}_{loc.get('shape_index')}_{loc.get('row_index')}_{loc.get('col_index')}"
            elif loc_type == "pptx_speaker_notes":
                key = f"pptx_speaker_notes_{loc.get('slide_index')}_{loc.get('paragraph_index')}"
            elif loc_type == "pdf_text_block":
                key = f"pdf_text_block_{loc.get('page_index')}_{loc.get('block_no')}"
            else:
                key = f"seg_{seg.segment_index}"

            segments_by_loc[key] = text

        opts = options or {}
        if t == "docx":
            DocxRenderer.render(working_copy, output_path, segments_by_loc)
            # Process embedded images if translate_images option is enabled
            if opts.get("translate_images", False):
                ocr_mode = opts.get("ocr_mode", "paddleocr")
                try:
                    import docx
                    out_doc = docx.Document(str(output_path))
                    img_parts = [p for p in out_doc.part.related_parts.values() if "image" in str(p.partname).lower()]
                    if img_parts:
                        logger.info(f"Processing {len(img_parts)} embedded images with ImageTranslator (OCR Engine: {ocr_mode})...")
                        for idx, img_p in enumerate(img_parts):
                            try:
                                new_bytes = await image_translator.process_image(
                                    image_bytes=img_p._blob,
                                    src_lang=src_lang,
                                    tgt_lang=tgt_lang,
                                    ocr_engine=ocr_mode,
                                    provider=provider
                                )
                                if new_bytes and len(new_bytes) > 0:
                                    img_p._blob = new_bytes
                            except Exception as img_err:
                                logger.warning(f"Failed to translate image {idx + 1}: {img_err}")
                        out_doc.save(str(output_path))
                        logger.info(f"Saved DOCX with {len(img_parts)} processed images to {output_path}")
                except Exception as img_exc:
                    logger.warning(f"Failed to perform embedded image translation on DOCX: {img_exc}")
        elif t == "xlsx":
            await XlsxRenderer.render(
                working_copy,
                output_path,
                segments_by_loc,
                options=opts,
                src_lang=src_lang,
                tgt_lang=tgt_lang,
                provider=provider
            )
        elif t == "pptx":
            await PptxRenderer.render(
                working_copy,
                output_path,
                segments_by_loc,
                options=opts,
                src_lang=src_lang,
                tgt_lang=tgt_lang,
                provider=provider
            )
        elif t == "pdf":
            await PdfRenderer.render(
                working_copy,
                output_path,
                segments_by_loc,
                raw_segs,
                options=opts,
                src_lang=src_lang,
                tgt_lang=tgt_lang,
                provider=provider
            )

job_manager = JobManager()
