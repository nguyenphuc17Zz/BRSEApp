import io
import os
from pathlib import Path
from typing import Dict, Any, Optional
import pptx
from pptx.util import Pt
from pptx.oxml.ns import qn
from PIL import Image

from app.core.logging import logger
from app.documents.ocr.image_translator import image_translator


class PptxRenderer:
    """Renders translated text into PowerPoint (.pptx) presentations with overflow handling and in-place image translation."""

    @classmethod
    async def render(
        cls,
        working_path: Path,
        output_path: Path,
        segments_by_loc: Dict[str, str],
        options: Optional[Dict[str, Any]] = None,
        src_lang: str = "ja",
        tgt_lang: str = "vi",
        provider: Any = None
    ):
        prs = pptx.Presentation(str(working_path))
        opts = options or {}
        translate_images = opts.get("translate_images", True)
        ocr_mode = opts.get("ocr_mode", "auto")
        is_vi = (tgt_lang or "").lower() in ("vi", "vietnamese")
        font_name = "Arial" if is_vi else "MS Gothic"

        for slide_idx, slide in enumerate(prs.slides):
            # 1. Shapes & Text frames
            for shape_idx, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                        loc_key = f"pptx_shape_text_{slide_idx}_{shape_idx}_{p_idx}"
                        if loc_key in segments_by_loc:
                            new_text = segments_by_loc[loc_key]
                            if is_vi:
                                new_text = new_text.replace("・", "• ")

                            old_len = len(paragraph.text)
                            paragraph.text = new_text

                            # Assign language-appropriate font to runs
                            for run in paragraph.runs:
                                run.font.name = font_name

                            # Safe font scaling if text expanded significantly
                            if len(new_text) > old_len * 1.3 and paragraph.font.size:
                                current_pt = paragraph.font.size.pt
                                new_pt = max(9.0, current_pt * 0.9)
                                paragraph.font.size = Pt(new_pt)

                # Tables in slide
                elif shape.has_table:
                    for r_idx, row in enumerate(shape.table.rows):
                        for c_idx, cell in enumerate(row.cells):
                            loc_key = f"pptx_table_cell_{slide_idx}_{shape_idx}_{r_idx}_{c_idx}"
                            if loc_key in segments_by_loc:
                                new_text = segments_by_loc[loc_key]
                                if is_vi:
                                    new_text = new_text.replace("・", "• ")
                                cell.text = new_text
                                for p in cell.text_frame.paragraphs:
                                    for run in p.runs:
                                        run.font.name = font_name

            # 2. Speaker Notes
            if slide.has_notes_slide:
                notes_tf = slide.notes_slide.notes_text_frame
                if notes_tf:
                    for p_idx, paragraph in enumerate(notes_tf.paragraphs):
                        loc_key = f"pptx_speaker_notes_{slide_idx}_{p_idx}"
                        if loc_key in segments_by_loc:
                            new_text = segments_by_loc[loc_key]
                            if is_vi:
                                new_text = new_text.replace("・", "• ")
                            paragraph.text = new_text
                            for run in paragraph.runs:
                                run.font.name = font_name

        # 3. In-place Image Translation across slides
        if translate_images:
            processed_rIds = set()
            for slide_idx, slide in enumerate(prs.slides):
                for shape in slide.shapes:
                    blips = []
                    try:
                        blips = shape._element.xpath('.//a:blip')
                    except Exception:
                        pass

                    for blip in blips:
                        rId = getattr(blip, "rEmbed", None) or blip.get(qn('r:embed'))
                        if not rId or rId in processed_rIds:
                            continue
                        processed_rIds.add(rId)

                        try:
                            img_part = slide.part.related_part(rId)
                            if not img_part or not hasattr(img_part, "blob"):
                                continue

                            raw_bytes = img_part.blob
                            pil_temp = Image.open(io.BytesIO(raw_bytes))
                            if pil_temp.width < 60 or pil_temp.height < 60:
                                continue

                            logger.info(
                                f"Processing PPTX embedded image on slide {slide_idx + 1} "
                                f"({pil_temp.width}x{pil_temp.height}, OCR: {ocr_mode})..."
                            )
                            new_bytes = await image_translator.process_image(
                                image_bytes=raw_bytes,
                                src_lang=src_lang,
                                tgt_lang=tgt_lang,
                                ocr_engine=ocr_mode,
                                provider=provider
                            )
                            if new_bytes and len(new_bytes) > 0 and new_bytes != raw_bytes:
                                img_part._blob = new_bytes
                                logger.info(
                                    f"Successfully replaced image {rId} on slide {slide_idx + 1} "
                                    f"with translated version."
                                )
                        except Exception as img_err:
                            logger.warning(
                                f"Failed to translate embedded image {rId} on slide {slide_idx + 1}: {img_err}"
                            )

        prs.save(str(output_path))
        logger.info(f"PPTX presentation successfully rendered to {output_path}")
