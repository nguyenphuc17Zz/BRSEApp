from pathlib import Path
from typing import Dict, Any, Optional
import re
import json
import openpyxl
from app.core.logging import logger
from app.documents.ocr.image_translator import image_translator
from app.documents.parsers.formula_helper import extract_formula_strings, replace_formula_strings

class XlsxRenderer:
    """Renders translated text into Excel (.xlsx) workbook, translating formula strings and inpainting embedded images."""

    @classmethod
    async def render(
        cls,
        working_path: Path,
        output_path: Path,
        segments_by_loc: Dict[str, str],
        options: Optional[Dict[str, Any]] = None,
        src_lang: str = "ja",
        tgt_lang: str = "vi",
        provider: Any = None
    ):
        wb = openpyxl.load_workbook(str(working_path), data_only=False)
        opts = options or {}
        selected_sheets = opts.get("selected_sheets")
        translate_images = opts.get("translate_images", False)
        ocr_mode = opts.get("ocr_mode", "paddleocr")

        for sheet_name in wb.sheetnames:
            if selected_sheets and sheet_name not in selected_sheets:
                continue

            ws = wb[sheet_name]
            # 1. Update text cells and formula string literals
            for row_idx in range(1, ws.max_row + 1):
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    val = cell.value

                    # Handle formula cells: preserve formula structure, translate natural language strings
                    if isinstance(val, str) and val.startswith("="):
                        formula_strings = extract_formula_strings(val)
                        if formula_strings:
                            replacements = {}
                            for item in formula_strings:
                                str_key = f"excel_formula_string_{sheet_name}_{row_idx}_{col_idx}_{item['str_index']}"
                                if str_key in segments_by_loc:
                                    replacements[item["str_index"]] = segments_by_loc[str_key]
                            if replacements:
                                new_val = replace_formula_strings(val, replacements)
                                cell.value = new_val
                        continue

                    # Handle regular text cells
                    loc_key = f"excel_cell_{sheet_name}_{row_idx}_{col_idx}"
                    if loc_key in segments_by_loc:
                        cell.value = segments_by_loc[loc_key]

            # 2. Process embedded images in this worksheet
            if translate_images and hasattr(ws, "_images") and ws._images:
                logger.info(f"Processing {len(ws._images)} embedded images in sheet '{sheet_name}' (OCR Engine: {ocr_mode})...")
                for img_idx, img in enumerate(ws._images):
                    try:
                        raw_bytes = img._data() if hasattr(img, "_data") and callable(img._data) else None
                        if raw_bytes:
                            new_bytes = await image_translator.process_image(
                                image_bytes=raw_bytes,
                                src_lang=src_lang,
                                tgt_lang=tgt_lang,
                                ocr_engine=ocr_mode,
                                provider=provider
                            )
                            if new_bytes and len(new_bytes) > 0:
                                img._data = lambda nb=new_bytes: nb
                                try:
                                    import io
                                    from PIL import Image as PILImage
                                    img.ref = PILImage.open(io.BytesIO(new_bytes))
                                except Exception:
                                    pass
                    except Exception as img_err:
                        logger.warning(f"Failed to translate embedded image {img_idx + 1} in sheet '{sheet_name}': {img_err}")

        # 3. Rename sheets if translate_sheet_names is enabled
        translate_sheet_names = opts.get("translate_sheet_names", True)
        if translate_sheet_names and wb.sheetnames:
            sheet_title_map = {}
            if provider:
                items = [{"id": i, "text": name} for i, name in enumerate(wb.sheetnames)]
                prompt = f"""Translate the following Excel sheet names from {src_lang} to {tgt_lang}. Keep them concise (max 30 chars).
JSON:
{json.dumps(items, ensure_ascii=False)}

Schema: {{"translations": [{{"id": 0, "translated": "..."}}]}}"""
                try:
                    from app.engine.pipeline import clean_json_response
                    resp = await provider.generate(
                        prompt=prompt,
                        system_instruction=f"Translate spreadsheet tab names concisely from {src_lang} to {tgt_lang}.",
                        temperature=0.1,
                        json_mode=True
                    )
                    data = clean_json_response(resp.text)
                    for t in data.get("translations", []):
                        raw_id = t.get("id")
                        val = t.get("translated", "")
                        if val:
                            sheet_title_map[raw_id] = val
                            try:
                                sheet_title_map[int(raw_id)] = val
                                sheet_title_map[str(raw_id)] = val
                            except (ValueError, TypeError):
                                pass
                except Exception as sheet_err:
                    logger.warning(f"Failed to translate Excel sheet names via AI: {sheet_err}")

            used_names = set()
            for idx, old_name in enumerate(list(wb.sheetnames)):
                new_title = sheet_title_map.get(idx) or sheet_title_map.get(str(idx))
                if new_title and new_title.strip():
                    clean_name = re.sub(r'[\\/?*\[\]:]', '_', new_title.strip())[:31]
                    final_name = clean_name
                    counter = 1
                    while final_name in used_names or (final_name in wb.sheetnames and final_name != old_name):
                        suffix = f"_{counter}"
                        final_name = f"{clean_name[:31-len(suffix)]}{suffix}"
                        counter += 1
                    try:
                        wb[old_name].title = final_name
                        used_names.add(final_name)
                    except Exception as ren_err:
                        logger.warning(f"Could not rename Excel sheet '{old_name}' to '{final_name}': {ren_err}")
                else:
                    used_names.add(old_name)

        wb.save(str(output_path))
        logger.info(f"XLSX workbook successfully rendered to {output_path}")
