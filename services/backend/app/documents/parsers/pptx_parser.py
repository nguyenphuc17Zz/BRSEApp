from pathlib import Path
from typing import Dict, List, Tuple, Any, Optional
import pptx
from app.documents.parsers.base import DocumentParser
from app.documents.segmenter import ParsedSegment, TokenProtector

class PptxParser(DocumentParser):
    """Parses PowerPoint (.pptx) presentations preserving slide structure, shapes, tables, and notes."""

    def parse(self, file_path: Path, options: Optional[Dict[str, Any]] = None) -> Tuple[List[ParsedSegment], Dict[str, Any]]:
        prs = pptx.Presentation(str(file_path))
        opts = options or {}
        translate_notes = opts.get("translate_notes", True)

        segments: List[ParsedSegment] = []
        segment_index = 0

        for slide_idx, slide in enumerate(prs.slides):
            # Extract slide title if present to provide structured context
            slide_title = f"Slide {slide_idx + 1}"
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()

            # 1. Parse Shapes & Text Frames
            for shape_idx, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                        text = paragraph.text.strip()
                        if not text:
                            continue

                        protected_text, token_map = TokenProtector.protect_tokens(text)
                        segments.append(ParsedSegment(
                            segment_index=segment_index,
                            source_text=protected_text,
                            location={
                                "type": "pptx_shape_text",
                                "slide_index": slide_idx,
                                "shape_index": shape_idx,
                                "paragraph_index": p_idx
                            },
                            protected_tokens=token_map,
                            context_hint=f"Slide {slide_idx + 1}: {slide_title}",
                            formatting_meta={
                                "is_title": shape == slide.shapes.title,
                                "font_size": paragraph.font.size.pt if paragraph.font and paragraph.font.size else None
                            }
                        ))
                        segment_index += 1

                # Tables inside slides
                elif shape.has_table:
                    for r_idx, row in enumerate(shape.table.rows):
                        for c_idx, cell in enumerate(row.cells):
                            text = cell.text.strip()
                            if not text:
                                continue

                            protected_text, token_map = TokenProtector.protect_tokens(text)
                            segments.append(ParsedSegment(
                                segment_index=segment_index,
                                source_text=protected_text,
                                location={
                                    "type": "pptx_table_cell",
                                    "slide_index": slide_idx,
                                    "shape_index": shape_idx,
                                    "row_index": r_idx,
                                    "col_index": c_idx
                                },
                                protected_tokens=token_map,
                                context_hint=f"Slide {slide_idx + 1} Table - Row {r_idx + 1}, Col {c_idx + 1}",
                                formatting_meta={"is_table": True}
                            ))
                            segment_index += 1

            # 2. Speaker Notes
            if translate_notes and slide.has_notes_slide:
                notes_text_frame = slide.notes_slide.notes_text_frame
                if notes_text_frame and notes_text_frame.text.strip():
                    for p_idx, paragraph in enumerate(notes_text_frame.paragraphs):
                        text = paragraph.text.strip()
                        if not text:
                            continue

                        protected_text, token_map = TokenProtector.protect_tokens(text)
                        segments.append(ParsedSegment(
                            segment_index=segment_index,
                            source_text=protected_text,
                            location={
                                "type": "pptx_speaker_notes",
                                "slide_index": slide_idx,
                                "paragraph_index": p_idx
                            },
                            protected_tokens=token_map,
                            context_hint=f"Slide {slide_idx + 1} Speaker Notes",
                            formatting_meta={"is_speaker_notes": True}
                        ))
                        segment_index += 1

        metadata = {
            "unit_count": len(prs.slides),
            "unit_label": "slides",
            "slide_count": len(prs.slides),
            "total_segments": len(segments)
        }
        return segments, metadata
