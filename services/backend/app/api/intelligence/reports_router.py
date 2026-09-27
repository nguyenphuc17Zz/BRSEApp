import os
import uuid
import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any
from pydantic import BaseModel, Field
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import FileResponse
from sqlalchemy import select, desc, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.logging import logger
from app.intelligence.models import ReportRecord
from app.intelligence.reports.report_service import ReportService
from app.intelligence.reports.slide_report_service import SlideReportService

router = APIRouter(prefix="/api/intelligence/reports", tags=["Reports"])

TEMPLATE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../../data/templates/reports"))
EXPORTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../../data/exports/reports"))

os.makedirs(TEMPLATE_DIR, exist_ok=True)
os.makedirs(EXPORTS_DIR, exist_ok=True)

class GenerateReportRequest(BaseModel):
    project_id: str
    report_type: str = "client_nippo" # client_nippo, client_shuho, internal_standup
    is_auto_harvest: bool = False
    manual_input_raw: Optional[str] = ""
    target_language: str = "ja" # ja, vi, bilingual
    sender_name: str = "BrSE / Offshore Lead"
    recipient_name: str = "お客様 (Client PM / Tech Lead)"
    selected_item_ids: List[str] = Field(default_factory=list)
    additional_notes: Optional[str] = ""
    report_date: Optional[str] = None
    template_file_path: Optional[str] = None
    custom_text_template: Optional[str] = ""
    provider: str = "auto"
    model: Optional[str] = None

@router.get("/harvest")
async def harvest_report_data(
    project_id: str = Query(..., description="Project ID"),
    report_type: str = Query("client_nippo", description="Report type"),
    date: Optional[str] = Query(None, description="Report date YYYY-MM-DD"),
    db: AsyncSession = Depends(get_db)
):
    """Harvests project tasks, blockers, and questions for manual review before generation."""
    try:
        data = await ReportService.harvest_report_data(db, project_id, report_type, date)
        return data
    except Exception as e:
        logger.error(f"Error harvesting report data: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/upload-template")
async def upload_slide_template(
    file: UploadFile = File(...)
):
    """Uploads a presentation (.pptx), document (.docx), or text template (.txt, .md)."""
    ext = Path(file.filename).suffix.lower()
    if ext not in [".pptx", ".docx", ".xlsx", ".txt", ".md"]:
        raise HTTPException(status_code=400, detail="Only .pptx, .docx, .xlsx, .txt, and .md templates are supported.")

    file_id = str(uuid.uuid4())[:8]
    safe_name = f"{file_id}_{file.filename}"
    save_path = os.path.join(TEMPLATE_DIR, safe_name)

    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    analysis = None
    raw_text = None
    if ext == ".pptx":
        try:
            analysis = SlideReportService.analyze_slide_template(save_path)
        except Exception as e:
            logger.warning(f"PPTX analysis warning: {e}")
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(save_path)
            analysis = {
                "file_type": "docx",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables)
            }
        except Exception as e:
            logger.warning(f"DOCX analysis warning: {e}")
    elif ext in [".txt", ".md"]:
        try:
            raw_text = content.decode("utf-8", errors="replace")
            analysis = {
                "file_type": "text",
                "char_count": len(raw_text)
            }
        except Exception as e:
            logger.warning(f"Text template decode warning: {e}")

    return {
        "status": "uploaded",
        "filename": file.filename,
        "template_file_path": save_path,
        "analysis": analysis,
        "raw_text": raw_text
    }

@router.post("/generate")
async def generate_report(
    req: GenerateReportRequest,
    db: AsyncSession = Depends(get_db)
):
    """Generates a professional business report across 4 output channels (HTML, Chatwork, Slack, Markdown)."""
    try:
        result = await ReportService.generate_report(req.model_dump(), db)
        return result
    except Exception as e:
        logger.error(f"Error generating report: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/history")
async def get_report_history(
    project_id: str = Query(..., description="Project ID"),
    limit: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db)
):
    """Returns past generated reports for the project."""
    query = (
        select(ReportRecord)
        .where(ReportRecord.project_id == project_id)
        .order_by(desc(ReportRecord.created_at))
        .limit(limit)
    )
    records = (await db.execute(query)).scalars().all()
    return [
        {
            "id": r.id,
            "project_id": r.project_id,
            "title": r.title,
            "report_type": r.report_type,
            "report_date": r.report_date,
            "recipient_name": r.recipient_name,
            "sender_name": r.sender_name,
            "is_auto_harvest": r.is_auto_harvest,
            "target_language": r.target_language,
            "template_file_path": r.template_file_path,
            "content_markdown": r.content_markdown,
            "content_html": r.content_html,
            "content_chatwork": r.content_chatwork,
            "content_slack": r.content_slack,
            "content_vietnamese_preview": r.content_vietnamese_preview,
            "created_at": r.created_at.isoformat() if r.created_at else ""
        }
        for r in records
    ]

@router.get("/{report_id}")
async def get_single_report(
    report_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Retrieves a single report by ID."""
    record = (await db.execute(select(ReportRecord).where(ReportRecord.id == report_id))).scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")
    return {
        "id": record.id,
        "project_id": record.project_id,
        "title": record.title,
        "report_type": record.report_type,
        "report_date": record.report_date,
        "recipient_name": record.recipient_name,
        "sender_name": record.sender_name,
        "is_auto_harvest": record.is_auto_harvest,
        "target_language": record.target_language,
        "template_file_path": record.template_file_path,
        "content_markdown": record.content_markdown,
        "content_html": record.content_html,
        "content_chatwork": record.content_chatwork,
        "content_slack": record.content_slack,
        "content_vietnamese_preview": record.content_vietnamese_preview,
        "created_at": record.created_at.isoformat() if record.created_at else ""
    }

@router.delete("/{report_id}")
async def delete_report(
    report_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Deletes a report by ID."""
    record = (await db.execute(select(ReportRecord).where(ReportRecord.id == report_id))).scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")
    await db.delete(record)
    await db.commit()
    return {"status": "deleted", "id": report_id}

@router.get("/{report_id}/export-docx")
async def export_report_docx(
    report_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Generates and downloads Microsoft Word (.docx) document."""
    record = (await db.execute(select(ReportRecord).where(ReportRecord.id == report_id))).scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")

    safe_title = "".join(c for c in record.title if c.isalnum() or c in " _-").strip() or "Report"
    filename = f"{safe_title}_{record.report_date}.docx"
    output_path = os.path.join(EXPORTS_DIR, filename)

    ReportService.export_report_docx(record, output_path)
    return FileResponse(
        path=output_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

@router.get("/{report_id}/export-xlsx")
async def export_report_xlsx(
    report_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Generates and downloads Microsoft Excel (.xlsx) workbook."""
    record = (await db.execute(select(ReportRecord).where(ReportRecord.id == report_id))).scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")

    safe_title = "".join(c for c in record.title if c.isalnum() or c in " _-").strip() or "Report"
    filename = f"{safe_title}_{record.report_date}.xlsx"
    output_path = os.path.join(EXPORTS_DIR, filename)

    ReportService.export_report_xlsx(record, output_path)
    return FileResponse(
        path=output_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@router.get("/{report_id}/export-pptx")
async def export_report_pptx(
    report_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Generates or fills a PowerPoint (.pptx) presentation from the report data."""
    record = (await db.execute(select(ReportRecord).where(ReportRecord.id == report_id))).scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")

    safe_title = "".join(c for c in record.title if c.isalnum() or c in " _-").strip() or "Weekly_Slide"
    filename = f"{safe_title}_{record.report_date}.pptx"
    output_path = os.path.join(EXPORTS_DIR, filename)

    ReportService.export_report_pptx(record, output_path)
    return FileResponse(
        path=output_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
